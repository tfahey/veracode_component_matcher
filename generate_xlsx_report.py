#!/usr/bin/env python3
"""
Generate comprehensive SBOM/SCA analysis Excel report with all matching strategies.

This script automates the Excel generation process currently done manually,
using the VulnerabilityMatcher with all 5 matching strategies to achieve
higher match rates and provide complete traceability.
"""

import json
import logging
import argparse
from datetime import datetime
from typing import Optional
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from api_fetcher import VeracodeApiFetcher
from vulnerability_matcher import VulnerabilityMatcher

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# Excel styling constants
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MATCH_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
UNMATCH_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
STRATEGY_COLORS = {
    'sha1': PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    'purl': PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    'composite_key': PatternFill(start_color="E8F4F8", end_color="E8F4F8", fill_type="solid"),
    'library_id': PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    'fuzzy_name': PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid"),
}


class ExcelReportGenerator:
    """Generate SBOM/SCA analysis Excel workbooks"""

    def __init__(self, fetcher: Optional[VeracodeApiFetcher] = None, matcher: Optional[VulnerabilityMatcher] = None):
        self.fetcher = fetcher
        self.matcher = matcher or VulnerabilityMatcher()
        self.wb = openpyxl.Workbook()
        if self.wb.active is not None:
            self.wb.remove(self.wb.active)

    def _apply_header_style(self, ws, row: int, num_cols: int) -> None:
        """Apply header styling to a row"""
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _auto_adjust_columns(self, ws) -> None:
        """Auto-adjust column widths"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value or '')) > max_length:
                        max_length = len(str(cell.value or ''))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def add_sbom_sheet(self) -> None:
        """Add SBOM Components sheet"""
        logger.info("Adding SBOM Components sheet...")
        ws = self.wb.create_sheet("SBOM Components")

        headers = [
            'Application ID', 'App Name', 'Component Name', 'Component Version',
            'SHA1', 'SHA256', 'PURL', 'BOM-Ref', 'Match Found', 'Match Strategy',
            'Matched SCA Name', 'Matched SCA Version', 'CVEs', 'Vulnerabilities'
        ]

        ws.append(headers)
        self._apply_header_style(ws, 1, len(headers))

        # Add SBOM data with match results
        for sbom_comp in self.matcher.sbom_components.values():
            # Find if this component was matched
            match_result = next(
                (m for m in self.matcher.matched if m.sbom_component == sbom_comp),
                None
            )

            match_strategy = ''
            sca_name = ''
            sca_version = ''
            vuln_count = 0
            strategy_key = None

            if match_result and match_result.match_strategy:
                match_strategy = match_result.match_strategy.upper()
                strategy_key = match_result.match_strategy

            if match_result and match_result.sca_component:
                sca_name = match_result.sca_component.name
                sca_version = match_result.sca_component.version
                vuln_count = len(match_result.sca_component.vulnerabilities) if match_result.sca_component.vulnerabilities else 0

            row_data = [
                sbom_comp.vendor or '',  # Using vendor as app-level info
                '',  # App name - would need to track from original data
                sbom_comp.name,
                sbom_comp.version,
                sbom_comp.sha1 or '',
                '',  # SHA256 not currently tracked
                sbom_comp.purl or '',
                sbom_comp.bom_ref or '',
                'Yes' if match_result else 'No',
                match_strategy,
                sca_name,
                sca_version,
                '',  # CVEs from metadata
                vuln_count
            ]

            ws.append(row_data)

            # Color code matched rows
            if match_result and strategy_key:
                fill = STRATEGY_COLORS.get(strategy_key, MATCH_FILL)
                for cell in ws[ws.max_row]:
                    cell.fill = fill

        self._auto_adjust_columns(ws)
        logger.info(f"  Added {len(self.matcher.sbom_components)} SBOM components")

    def add_sca_sheet(self) -> None:
        """Add SCA Vulnerability Data sheet"""
        logger.info("Adding SCA Vulnerability Data sheet...")
        ws = self.wb.create_sheet("SCA Vuln Data")

        headers = [
            'Library Name', 'Library Version', 'Library ID', 'SHA1',
            'CVE ID', 'Severity', 'CWE ID', 'CVSS Score',
            'Matched to SBOM', 'SBOM Component Name', 'Vulnerabilities'
        ]

        ws.append(headers)
        self._apply_header_style(ws, 1, len(headers))

        # Add SCA data
        for sca_comp in self.matcher.sca_components.values():
            # Find if this was matched
            match_result = next(
                (m for m in self.matcher.matched if m.sca_component == sca_comp),
                None
            )

            vuln_cves = ''
            if sca_comp.vulnerabilities:
                cve_list = []
                for v in sca_comp.vulnerabilities[:3]:
                    if isinstance(v, dict) and 'cve_id' in v:
                        cve_list.append(str(v['cve_id']))
                vuln_cves = ', '.join(cve_list)

            sbom_name = ''
            strategy_key = None
            if match_result:
                if match_result.sbom_component:
                    sbom_name = match_result.sbom_component.name
                if match_result.match_strategy:
                    strategy_key = match_result.match_strategy

            row_data = [
                sca_comp.name,
                sca_comp.version,
                sca_comp.library_id or '',
                sca_comp.sha1 or '',
                vuln_cves,
                '',  # Severity would come from vulnerabilities
                '',  # CWE would come from vulnerabilities
                '',  # CVSS Score
                'Yes' if match_result else 'No',
                sbom_name,
                len(sca_comp.vulnerabilities) if sca_comp.vulnerabilities else 0
            ]

            ws.append(row_data)

            # Color code matched rows
            if match_result and strategy_key:
                fill = STRATEGY_COLORS.get(strategy_key, MATCH_FILL)
                for cell in ws[ws.max_row]:
                    cell.fill = fill

        self._auto_adjust_columns(ws)
        logger.info(f"  Added {len(self.matcher.sca_components)} SCA components")

    def add_matched_sheet(self) -> None:
        """Add Matched Components sheet with detailed mapping"""
        logger.info("Adding Matched Components sheet...")
        ws = self.wb.create_sheet("Matched Components")

        headers = [
            'SBOM Component', 'SBOM Version', 'SCA Component', 'SCA Version',
            'Match Strategy', 'Confidence', 'SHA1 Match', 'PURL Match',
            'Vulnerabilities', 'CVE Examples'
        ]

        ws.append(headers)
        self._apply_header_style(ws, 1, len(headers))

        for result in self.matcher.matched:
            cves = ''
            if result.sca_component and result.sca_component.vulnerabilities:
                cve_list = []
                for v in result.sca_component.vulnerabilities[:5]:
                    if isinstance(v, dict) and 'cve_id' in v:
                        cve_list.append(str(v['cve_id']))
                cves = ', '.join(cve_list)
                if len(result.sca_component.vulnerabilities) > 5:
                    cves += f" +{len(result.sca_component.vulnerabilities) - 5} more"

            sca_name = result.sca_component.name if result.sca_component else ''
            sca_version = result.sca_component.version if result.sca_component else ''
            strategy = result.match_strategy.upper() if result.match_strategy else ''
            vuln_count = len(result.sca_component.vulnerabilities) if result.sca_component and result.sca_component.vulnerabilities else 0

            sha1_match = 'No'
            if result.sbom_component.sha1 and result.sca_component and result.sca_component.sha1:
                sha1_match = 'Yes' if result.sbom_component.sha1 == result.sca_component.sha1 else 'No'

            purl_match = 'No'
            if result.sbom_component.purl and result.sca_component and result.sca_component.purl:
                purl_match = 'Yes' if result.sbom_component.purl == result.sca_component.purl else 'No'

            row_data = [
                result.sbom_component.name,
                result.sbom_component.version,
                sca_name,
                sca_version,
                strategy,
                f"{result.confidence:.0%}",
                sha1_match,
                purl_match,
                vuln_count,
                cves
            ]

            ws.append(row_data)

            # Color by strategy
            fill = STRATEGY_COLORS.get(result.match_strategy or '', MATCH_FILL)
            for cell in ws[ws.max_row]:
                cell.fill = fill

        self._auto_adjust_columns(ws)
        logger.info(f"  Added {len(self.matcher.matched)} matched components")

    def add_unmatched_sheet(self) -> None:
        """Add Unmatched Components sheet with gap analysis"""
        logger.info("Adding Unmatched Components sheet...")
        ws = self.wb.create_sheet("Unmatched Components")

        headers = [
            'Source', 'Component Name', 'Version', 'SHA1', 'Library ID',
            'Match Attempts', 'Reason for No Match'
        ]

        ws.append(headers)
        self._apply_header_style(ws, 1, len(headers))

        for component in self.matcher.unmatched:
            reason = self._analyze_match_failure(component)

            row_data = [
                component.source.upper(),
                component.name,
                component.version,
                component.sha1 or '',
                component.library_id or '',
                '5 (all)',  # We try all strategies
                reason
            ]

            ws.append(row_data)
            ws[ws.max_row].fill = UNMATCH_FILL

        self._auto_adjust_columns(ws)
        logger.info(f"  Added {len(self.matcher.unmatched)} unmatched components")

    def add_statistics_sheet(self) -> None:
        """Add Summary Statistics sheet"""
        logger.info("Adding Statistics sheet...")
        ws = self.wb.create_sheet("Statistics")

        # Title
        ws['A1'] = "SBOM & SCA Component Matching Analysis"
        ws['A1'].font = Font(bold=True, size=14)

        # Generated timestamp
        ws['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        # Statistics section
        row = 4
        ws[f'A{row}'] = "Matching Summary"
        ws[f'A{row}'].font = Font(bold=True, size=12)

        stats = self.matcher.get_statistics()
        row += 1

        stats_data = [
            ('Total SBOM Components', stats['total_sbom_components']),
            ('Total SCA Components', stats['total_sca_components']),
            ('Successfully Matched', stats['matched_components']),
            ('Unmatched Components', stats['unmatched_components']),
            ('Match Success Rate', f"{(stats['matched_components'] / max(stats['total_sbom_components'], 1) * 100):.1f}%"),
            ('Gap Percentage', f"{stats['match_gap_percentage']:.1f}%"),
        ]

        for label, value in stats_data:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            ws[f'B{row}'].alignment = Alignment(horizontal="right")
            row += 1

        # Matches by strategy
        row += 1
        ws[f'A{row}'] = "Matches by Strategy"
        ws[f'A{row}'].font = Font(bold=True, size=12)
        row += 1

        ws[f'A{row}'] = "Strategy"
        ws[f'B{row}'] = "Count"
        ws[f'C{row}'] = "Percentage"
        self._apply_header_style(ws, row, 3)

        for strategy, count in sorted(stats['matches_by_strategy'].items()):
            row += 1
            percentage = (count / max(stats['matched_components'], 1)) * 100
            ws[f'A{row}'] = strategy.upper()
            ws[f'B{row}'] = count
            ws[f'C{row}'] = f"{percentage:.1f}%"
            ws[f'A{row}'].fill = STRATEGY_COLORS.get(strategy, PatternFill())

        # Adjust columns
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15

    def _analyze_match_failure(self, component) -> str:
        """Analyze why a component didn't match"""
        if not component.sha1 and not component.library_id:
            return "No SHA1 or Library ID available"
        elif not component.name or not component.version:
            return "Missing name or version"
        else:
            return "No matching component found across all strategies"

    def generate(self, output_file: str = 'SBOM_Vuln_Analysis_Generated.xlsx') -> str:
        """Generate the complete Excel workbook"""
        logger.info(f"Generating Excel report to {output_file}...")

        try:
            self.add_sbom_sheet()
            self.add_sca_sheet()
            self.add_matched_sheet()
            self.add_unmatched_sheet()
            self.add_statistics_sheet()

            # Ensure output directory exists
            Path(output_file).parent.mkdir(parents=True, exist_ok=True)
            self.wb.save(output_file)
            logger.info(f"✓ Report successfully generated: {output_file}")
            return output_file

        except Exception as e:
            logger.error(f"Failed to generate Excel report: {e}")
            raise


def main():
    """Main entry point"""
    parser = argparse.ArgumentParser(description='Generate SBOM/SCA analysis Excel report')
    parser.add_argument('--output', default='SBOM_Vuln_Analysis_Generated.xlsx',
                        help='Output Excel file path')
    parser.add_argument('--sbom-file', help='Input SBOM JSON file')
    parser.add_argument('--sca-file', help='Input SCA JSON file')
    parser.add_argument('--fetch-live', action='store_true',
                        help='Fetch data live from Veracode APIs')
    parser.add_argument('--veracode_api_key_id', help='Veracode API Key ID')
    parser.add_argument('--veracode_api_key_secret', help='Veracode API Key Secret')

    args = parser.parse_args()

    try:
        # Initialize matcher
        matcher = VulnerabilityMatcher()

        # Load data
        if args.fetch_live:
            logger.info("Fetching data from Veracode APIs...")
            _fetcher = VeracodeApiFetcher(api_id=args.veracode_api_key_id,
                                         api_key=args.veracode_api_key_secret)
            # Would need to implement full fetch in api_fetcher.py
            logger.warning("Live fetch not yet fully implemented")
        else:
            # Load from files
            if args.sbom_file and Path(args.sbom_file).exists():
                logger.info(f"Loading SBOM from {args.sbom_file}...")
                with open(args.sbom_file) as f:
                    sbom_data = json.load(f)
                    matcher.load_sbom_data(sbom_data)

            if args.sca_file and Path(args.sca_file).exists():
                logger.info(f"Loading SCA from {args.sca_file}...")
                with open(args.sca_file) as f:
                    sca_data = json.load(f)
                    matcher.load_sca_data(sca_data)

        # Match components
        logger.info("Matching components...")
        matcher.match_all()

        # Generate report
        generator = ExcelReportGenerator(matcher=matcher)
        output_file = generator.generate(args.output)
        print(f"\n✓ Report generated: {output_file}")

    except Exception as e:
        logger.error(f"Error: {e}")
        raise


if __name__ == '__main__':
    main()
